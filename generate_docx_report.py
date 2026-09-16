"""
Generates the official Laboratory Record document (.docx) according to the
Siksha 'O' Anusandhan (Deemed to be University) template for CSE 4192: Machine Learning Projects with Python.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_PATH = os.path.join(BASE_DIR, "Laboratory_Record_CSE4192.docx")
LOGO_PATH = os.path.join(BASE_DIR, "soa_logo.png")
VIZ_DIR = os.path.join(BASE_DIR, "visualizations")

def set_cell_background(cell, fill_hex):
    """Sets cell background color in docx table."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding for table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx_report():
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # PAGE 1: COVER PAGE (Exact SOA University Template)
    # -------------------------------------------------------------
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_before = Pt(10)
    p_uni.paragraph_format.space_after = Pt(2)
    run_uni1 = p_uni.add_run("SIKSHA ‘O’ ANUSANDHAN\n")
    run_uni1.font.name = 'Times New Roman'
    run_uni1.font.size = Pt(16)
    run_uni1.font.bold = True
    run_uni1.font.color.rgb = RGBColor(15, 23, 42)
    
    run_uni2 = p_uni.add_run("(DEEMED TO BE UNIVERSITY)\n")
    run_uni2.font.name = 'Times New Roman'
    run_uni2.font.size = Pt(13)
    run_uni2.font.bold = True
    run_uni2.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph() # spacing

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(8)
    p_meta.paragraph_format.space_after = Pt(14)
    run_meta1 = p_meta.add_run("Admission Batch: ")
    run_meta1.font.bold = True
    p_meta.add_run("2023 – 2027                                  ")
    run_meta2 = p_meta.add_run("Session: ")
    run_meta2.font.bold = True
    p_meta.add_run("2025 – 2026")

    doc.add_paragraph() # spacing

    p_rec = doc.add_paragraph()
    p_rec.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rec.paragraph_format.space_before = Pt(10)
    p_rec.paragraph_format.space_after = Pt(4)
    r_rec = p_rec.add_run("Laboratory Record\n")
    r_rec.font.name = 'Times New Roman'
    r_rec.font.size = Pt(14)
    r_rec.font.bold = True

    p_subj = doc.add_paragraph()
    p_subj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subj.paragraph_format.space_after = Pt(20)
    r_subj = p_subj.add_run("Machine Learning Projects with Python\n(CSE 4192)")
    r_subj.font.name = 'Times New Roman'
    r_subj.font.size = Pt(16)
    r_subj.font.bold = True
    r_subj.font.color.rgb = RGBColor(2, 132, 199)

    p_subm = doc.add_paragraph()
    p_subm.paragraph_format.space_before = Pt(10)
    p_subm.paragraph_format.space_after = Pt(4)
    r_subm = p_subm.add_run("Submitted by")
    r_subm.font.name = 'Times New Roman'
    r_subm.font.size = Pt(12)
    r_subm.font.bold = True

    p_names = doc.add_paragraph()
    p_names.paragraph_format.left_indent = Inches(0.5)
    p_names.paragraph_format.line_spacing = 1.3
    members = [
        "1. Name: Tribhuwan Singh (Regd. No.: 2341019538)",
        "2. Name: Surajit Sahoo (Regd. No.: 2341019165)",
        "3. Name: Anwesha Srichandan (Regd. No.: 2341019594)",
        "4. Name: Priti Rani Maity (Regd. No.: 2341013065)"
    ]
    for m in members:
        r = p_names.add_run(m + "\n")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = True

    doc.add_paragraph() # spacing

    # University Logo
    if os.path.exists(LOGO_PATH):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(8)
        p_logo.paragraph_format.space_after = Pt(16)
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.8))

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(14)
    p_dept.paragraph_format.space_after = Pt(2)
    r_d1 = p_dept.add_run("Centre for Artificial Intelligence & Machine Learning\n")
    r_d1.font.name = 'Times New Roman'
    r_d1.font.size = Pt(12)
    r_d1.font.bold = True

    r_d2 = p_dept.add_run("Faculty of Engineering & Technology (ITER)\n")
    r_d2.font.name = 'Times New Roman'
    r_d2.font.size = Pt(11.5)
    r_d2.font.bold = True

    r_d3 = p_dept.add_run("Jagamohan Nagar, Jagamara, Bhubaneswar, Odisha – 751030")
    r_d3.font.name = 'Times New Roman'
    r_d3.font.size = Pt(10)
    r_d3.font.italic = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # Helper Functions for Sections & Content
    # -------------------------------------------------------------
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(2, 132, 199)
        return p

    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10.5)
        return p

    # -------------------------------------------------------------
    # SECTION 1: INTRODUCTION
    # -------------------------------------------------------------
    add_h1("1 Introduction")
    
    add_h2("1.1 Problem Statement")
    add_body(
        "Diabetes mellitus is a progressive metabolic disorder characterized by chronic hyperglycemia resulting from defects "
        "in insulin secretion, insulin action, or both. Early identification of high-risk individuals is crucial to avert long-term "
        "retinopathy, nephropathy, and cardiovascular complications. However, conventional clinical diagnosis relies on static univariate "
        "thresholds (e.g., fasting plasma glucose >= 126 mg/dL) that fail to capture non-linear metabolic interactions. Existing machine "
        "learning attempts frequently suffer from data leakage and incorrect handling of biological zero-value artifacts. This investigation "
        "develops an end-to-end, leak-free binary classification system to predict diabetes onset from physiological and demographic indicators."
    )

    add_h2("1.2 Aim")
    add_body(
        "The primary aim of this laboratory project is to design, implement, optimize, benchmark, and deploy an authoritative "
        "Multilayer Perceptron (MLP) neural network capable of predicting 5-year diabetes onset probability from patient biometric markers "
        "using the Pima Indians Diabetes Database, comparing its predictive performance against standard machine learning baselines."
    )

    add_h2("1.3 Objectives")
    add_body("The specific SMART research objectives of this project are:")
    objectives_list = [
        "1. Conduct exhaustive Exploratory Data Analysis (EDA) quantifying target class imbalance (65.1% negative vs. 34.9% positive), skewness, and pairwise correlation structures.",
        "2. Implement a leak-free preprocessing pipeline that detects biologically implausible zeros in physiological continuous variables and imputes them using training-fitted medians.",
        "3. Design 16 domain-engineered features (expanding 8 raw features to 24 processed dimensions) incorporating WHO BMI bins, ADA glycemic categories, metabolic ratios, and non-linear log transforms.",
        "4. Develop and evaluate 7 baseline machine learning classifiers (Logistic Regression, KNN, SVM, Decision Tree, Random Forest, Gradient Boosting, Baseline MLP).",
        "5. Systematically optimize the Multilayer Perceptron via Grid Search across hidden layer topologies, L2 penalties, learning rates, and batch sizes.",
        "6. Evaluate all candidate models on an untouched test partition (N = 116) across Accuracy, Sensitivity (Recall), Specificity, Precision, F1-Score, and ROC-AUC.",
        "7. Deploy the production pipeline to an interactive Streamlit Cloud dashboard supporting single-patient triage, counterfactual sensitivity simulation, and batch cohort screening."
    ]
    for obj in objectives_list:
        p = doc.add_paragraph(obj)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)

    # -------------------------------------------------------------
    # SECTION 2: DATASET DESCRIPTION
    # -------------------------------------------------------------
    add_h1("2 Dataset Description")

    add_h2("2.1 Dataset Source")
    add_body(
        "The investigation utilizes the standard Pima Indians Diabetes Database collected by the National Institute of "
        "Diabetes and Digestive and Kidney Diseases (NIDDK). The dataset comprises records of female patients of Pima Indian heritage aged 21 years and older."
    )

    add_h2("2.2 Dataset Features")
    add_body("The dataset contains 8 raw clinical and physiological input attributes:")
    raw_feats = [
        "• Pregnancies: Total count of completed pregnancies.",
        "• Glucose: Plasma glucose concentration 2 hours after a 75g oral glucose tolerance test (OGTT) in mg/dL.",
        "• BloodPressure: Diastolic blood pressure reading in mm Hg.",
        "• SkinThickness: Triceps skinfold caliper thickness in mm (proxy for subcutaneous adiposity).",
        "• Insulin: 2-Hour post-load serum insulin concentration in μU/mL.",
        "• BMI: Body Mass Index computed as weight in kg / (height in m)^2.",
        "• DiabetesPedigreeFunction (DPF): Genetic risk score synthesized from family diabetes history.",
        "• Age: Chronological patient age in years."
    ]
    for rf in raw_feats:
        p = doc.add_paragraph(rf)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)

    add_h2("2.3 Target Variable")
    add_body(
        "The target variable is Outcome, a binary classification label indicating whether the patient developed diabetes within "
        "five years of initial examination (Outcome = 1 for diabetic, Outcome = 0 for non-diabetic)."
    )

    add_h2("2.4 Dataset Statistics")
    add_body("The dataset contains N = 768 patient instances. Summary descriptive statistics are presented in Table 1.")

    # Table 1: Dataset Summary Table
    t1 = doc.add_table(rows=1, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t1.rows[0].cells
    hdr_titles = ["Feature", "Type", "Mean ± SD", "Min", "Max", "Zero Count (%)"]
    for i, title in enumerate(hdr_titles):
        hdr[i].paragraphs[0].text = title
        hdr[i].paragraphs[0].runs[0].font.bold = True
        hdr[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr[i], "0F172A")
        set_cell_margins(hdr[i])

    ds_rows = [
        ["Pregnancies", "Discrete", "3.85 ± 3.37", "0", "17", "111 (14.45%)*"],
        ["Glucose", "Continuous", "120.89 ± 31.97", "0 (44)", "199", "5 (0.65%)"],
        ["BloodPressure", "Continuous", "69.11 ± 19.36", "0 (24)", "122", "35 (4.56%)"],
        ["SkinThickness", "Continuous", "20.54 ± 15.95", "0 (7)", "99", "227 (29.56%)"],
        ["Insulin", "Continuous", "79.80 ± 115.24", "0 (14)", "846", "374 (48.70%)"],
        ["BMI", "Continuous", "31.99 ± 7.88", "0 (18.2)", "67.1", "11 (1.43%)"],
        ["DiabetesPedigree", "Continuous", "0.47 ± 0.33", "0.078", "2.42", "0 (0.00%)"],
        ["Age", "Discrete", "33.24 ± 11.76", "21", "81", "0 (0.00%)"],
        ["Outcome", "Binary", "0.35 ± 0.48", "0", "1", "500 (65.1%) [0]"]
    ]
    for row_data in ds_rows:
        row = t1.add_row().cells
        for i, val in enumerate(row_data):
            row[i].paragraphs[0].text = val
            row[i].paragraphs[0].runs[0].font.size = Pt(9)
            set_cell_margins(row[i])

    add_body("*Note: Zero values in Pregnancies represent nulliparous women (biologically valid). Zeros in Glucose, BloodPressure, SkinThickness, Insulin, and BMI represent unrecorded measurements.")

    # -------------------------------------------------------------
    # SECTION 3: EXPLORATORY DATA ANALYSIS
    # -------------------------------------------------------------
    add_h1("3 Exploratory Data Analysis")

    add_h2("3.1 Dataset Structure")
    add_body("The dataset contains 768 rows and 9 columns with zero explicit SQL NULL values, but extensive hidden biological zero values.")

    add_h2("3.2 Descriptive Statistics")
    add_body(
        "Continuous physiological attributes exhibit pronounced variance. Insulin ranges from 0 to 846 μU/mL with a high standard "
        "deviation (115.24 μU/mL). Mean BMI is 31.99 kg/m^2, indicating that the patient cohort predominantly falls into the WHO Obese Class I category."
    )

    add_h2("3.3 Missing Value Analysis")
    add_body(
        "A biological sanity audit revealed that 374 records (48.70%) have Insulin = 0, 227 records (29.56%) have SkinThickness = 0, "
        "35 records (4.56%) have BloodPressure = 0, 11 records (1.43%) have BMI = 0, and 5 records (0.65%) have Glucose = 0. "
        "A living human cannot have zero blood glucose or zero blood pressure; these are missing data artifacts requiring imputation."
    )

    add_h2("3.4 Class Distribution")
    add_body(
        "The target variable consists of 500 negative instances (Class 0, 65.10%) and 268 positive instances (Class 1, 34.90%). "
        "The moderate class imbalance ratio (~1.87 : 1) necessitates evaluation via Precision, Recall, and ROC-AUC in addition to raw Accuracy."
    )

    add_h2("3.5 Univariate Analysis")
    add_body(
        "Insulin (skewness = 2.27) and DiabetesPedigreeFunction (skewness = 1.92) show severe positive right-skewness. "
        "Glucose and Blood Pressure approximate normal distributions once zero artifacts are removed."
    )

    add_h2("3.6 Bivariate Analysis")
    add_body(
        "Bivariate boxplot analysis demonstrates that diabetic patients exhibit statistically significant higher median glucose "
        "(140 vs 107 mg/dL), higher BMI (34.3 vs 30.1 kg/m^2), and older median age (36 vs 27 years)."
    )

    add_h2("3.7 Correlation Analysis")
    add_body(
        "Glucose exhibits the strongest linear correlation with diabetes outcome (r = 0.49), followed by BMI (r = 0.31), "
        "Age (r = 0.24), and Insulin (r = 0.21). Notable pairwise collinearity exists between Age and Pregnancies (r = 0.54) and SkinThickness and BMI (r = 0.65)."
    )

    add_h2("3.8 Outlier Analysis")
    add_body(
        "Interquartile Range (IQR) analysis identified physiological extremes (e.g., Insulin > 400 μU/mL, BMI > 50 kg/m^2). "
        "These were retained because extreme values represent genuine severe metabolic dysregulation rather than measurement errors."
    )

    # -------------------------------------------------------------
    # SECTION 4: DATA PREPROCESSING
    # -------------------------------------------------------------
    add_h1("4 Data Preprocessing")

    add_h2("4.1 Data Cleaning")
    add_body("Non-numeric entries and formatting anomalies were verified; column schemas were standardized to float64/int64.")

    add_h2("4.2 Missing Value Handling")
    add_body(
        "Zeros in continuous variables (Glucose, BloodPressure, SkinThickness, Insulin, BMI) were replaced with NaN. "
        "To strictly eliminate data leakage, median imputation statistics were fitted solely on the training partition (N = 537): "
        "Glucose = 117.0 mg/dL, BloodPressure = 72.0 mm Hg, SkinThickness = 29.0 mm, Insulin = 125.0 μU/mL, BMI = 32.0 kg/m^2."
    )

    add_h2("4.3 Outlier Handling")
    add_body("Non-linear log1p transformations were applied to skewed variables (Insulin, DPF) to stabilize numerical gradients without truncating data.")

    add_h2("4.4 Feature Engineering")
    add_body("16 domain features were engineered from imputed measurements, expanding the feature space from 8 to 24 dimensions:")
    eng_feats_desc = [
        "1. WHO Adiposity Bins (4 binary indicators): BMI_Underweight, BMI_Normal, BMI_Overweight, BMI_Obese.",
        "2. ADA Glycemic Stages (3 binary indicators): Glucose_Normal (<100), Glucose_Prediabetes (100-125), Glucose_Diabetes (>=126 mg/dL).",
        "3. Age Cohorts (3 binary indicators): Age_Young (<30), Age_Middle (30-50), Age_Senior (>50 years).",
        "4. Insulin_Glucose_Ratio: Insulin / (Glucose + 1e-5) — proxy for beta-cell compensation.",
        "5. Insulin_Resistance_Proxy: (Glucose * Insulin) / 405.0 — surrogate for systemic insulin resistance.",
        "6. Pregnancy_Age_Risk: Pregnancies / (Age + 1e-5) — normalizes parity against chronological age.",
        "7. BMI_Age_Interaction: BMI * Age — captures synergistic metabolic risk amplification.",
        "8. Log Transforms: Log_Insulin = ln(1 + Insulin) and Log_DPF = ln(1 + DPF) to compress tail distributions."
    ]
    for ef in eng_feats_desc:
        p = doc.add_paragraph(ef)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)

    add_h2("4.5 Feature Selection")
    add_body("All 24 features were retained to allow the Multilayer Perceptron to autonomously learn non-linear weight representations across linear, categorical, and interaction features.")

    add_h2("4.6 Train-Test Split")
    add_body("The dataset was partitioned using stratified sampling: 70% Training (N = 537), 15% Validation (N = 115), and 15% Untouched Test (N = 116), ensuring identical 34.9% positive class prevalence across partitions.")

    add_h2("4.7 Feature Scaling")
    add_body("StandardScaler was fitted on training data and applied to scale all 24 features to zero mean and unit variance (z = (x - μ_train) / σ_train).")

    # -------------------------------------------------------------
    # SECTION 5: METHODOLOGY
    # -------------------------------------------------------------
    add_h1("5 Methodology")

    add_h2("5.1 Overall Project Workflow")
    add_body("The project follows a rigorous machine learning lifecycle: Data Acquisition -> Zero Detection -> Training-Fitted Imputation -> Feature Engineering -> Stratified Partitioning -> Baseline Benchmarking -> MLP Hyperparameter Grid Search -> Final Retraining -> Untouched Test Evaluation -> Web Deployment.")

    add_h2("5.2 Model Development Strategy")
    add_body("Establish empirical baselines across multiple algorithmic families (linear, instance-based, margin-based, decision tree, tree ensemble) before optimizing the Multilayer Perceptron.")

    add_h2("5.3 Baseline Models")
    add_body("Six standard classifiers were trained under identical leak-free preprocessing: Logistic Regression, K-Nearest Neighbors, Support Vector Machine (RBF), Decision Tree (CART), Random Forest (100 trees), and Gradient Boosting.")

    add_h2("5.4 Multilayer Perceptron Architecture")
    add_body(
        "A feed-forward artificial neural network with an input layer of 24 neurons, two hidden layers (64 neurons and 32 neurons) "
        "equipped with ReLU activation functions, and a single sigmoid output neuron: z1 = W1*x + b1, a1 = ReLU(z1); z2 = W2*a1 + b2, a2 = ReLU(z2); z3 = W3*a2 + b3, y_hat = Sigmoid(z3)."
    )

    # -------------------------------------------------------------
    # SECTION 6: MODEL DEVELOPMENT
    # -------------------------------------------------------------
    add_h1("6 Model Development")
    add_body(
        "Each model family was instantiated with consistent random seeds (seed = 42):\n"
        "• 6.1 Logistic Regression: L2-regularized logistic regression with L-BFGS solver.\n"
        "• 6.2 K-Nearest Neighbors: Euclidean distance metric with k = 7 neighbors.\n"
        "• 6.3 Support Vector Machine: Radial Basis Function (RBF) kernel with C = 1.0 and gamma = 'scale'.\n"
        "• 6.4 Random Forest: 100 ensemble estimators with max_depth = 8 and Gini impurity criterion.\n"
        "• 6.5 Multilayer Perceptron: Scikit-Learn MLPClassifier with Adam optimizer, batch size 32, alpha = 0.001, early stopping with 10 epochs patience."
    )

    # -------------------------------------------------------------
    # SECTION 7: MODEL EVALUATION & SECTION 8: COMPARISON
    # -------------------------------------------------------------
    add_h1("7 Model Evaluation & Benchmark Results")
    add_body(
        "All models were evaluated on the untouched test partition (N = 116, 76 Negative, 40 Positive cases). "
        "Standard metrics were computed: Accuracy = (TP + TN) / N, Precision = TP / (TP + FP), Recall = TP / (TP + FN), "
        "Specificity = TN / (TN + FP), F1-Score = 2 * (Precision * Recall) / (Precision + Recall), and ROC-AUC."
    )

    # Table 2: Benchmark Comparison Table
    t2 = doc.add_table(rows=1, cols=9)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr2 = t2.rows[0].cells
    hdr2_titles = ["Algorithm", "Accuracy", "Sensitivity", "Specificity", "Precision", "F1", "ROC-AUC", "FN", "FP"]
    for i, title in enumerate(hdr2_titles):
        hdr2[i].paragraphs[0].text = title
        hdr2[i].paragraphs[0].runs[0].font.bold = True
        hdr2[i].paragraphs[0].runs[0].font.size = Pt(8.5)
        hdr2[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr2[i], "0F172A")
        set_cell_margins(hdr2[i], top=80, bottom=80, left=100, right=100)

    bench_rows = [
        ["Logistic Regression", "79.31%", "72.50%", "82.89%", "0.6905", "0.7073", "0.8444", "11", "13"],
        ["K-Nearest Neighbors", "81.90%", "77.50%", "84.21%", "0.7209", "0.7470", "0.8954", "9", "12"],
        ["Support Vector Machine", "83.62%", "77.50%", "86.84%", "0.7561", "0.7654", "0.8808", "9", "10"],
        ["Decision Tree", "87.93%", "82.50%", "90.79%", "0.8250", "0.8250", "0.9021", "7", "7"],
        ["Random Forest", "88.79%", "87.50%", "89.47%", "0.8140", "0.8434", "0.9431", "5", "8"],
        ["Gradient Boosting", "89.66%", "85.00%", "92.11%", "0.8500", "0.8500", "0.9579", "6", "6"],
        ["MLP (Baseline)", "81.03%", "65.00%", "89.47%", "0.7647", "0.7027", "0.8625", "14", "8"],
        ["MLP (Optimized)", "82.76%", "75.00%", "86.84%", "0.7500", "0.7500", "0.8681", "10", "10"]
    ]
    for row_data in bench_rows:
        row = t2.add_row().cells
        is_opt = "Optimized" in row_data[0]
        for i, val in enumerate(row_data):
            row[i].paragraphs[0].text = val
            row[i].paragraphs[0].runs[0].font.size = Pt(8)
            if is_opt:
                row[i].paragraphs[0].runs[0].font.bold = True
                set_cell_background(row[i], "E0F2FE")
            set_cell_margins(row[i], top=60, bottom=60, left=100, right=100)

    # -------------------------------------------------------------
    # SECTION 9: HYPERPARAMETER TUNING
    # -------------------------------------------------------------
    add_h1("9 Hyperparameter Tuning")
    add_body(
        "A systematic Grid Search was conducted across 8 candidate configurations over hidden layer architectures, "
        "L2 weight decay penalty (alpha), learning rates, and batch sizes. The optimal configuration: Hidden Topology: (64, 32), "
        "Activation: ReLU, Optimizer: Adam, alpha = 0.001, Initial Learning Rate = 0.001, Batch Size = 32 achieved highest validation ROC-AUC (0.8645)."
    )

    # -------------------------------------------------------------
    # SECTION 10 & 11: FINAL TRAINING & EVALUATION
    # -------------------------------------------------------------
    add_h1("10 Final Model Training & 11 Final Evaluation")
    add_body(
        "The optimal MLP architecture was retrained on combined Train + Validation data (N = 652) and evaluated on the untouched test partition (N = 116). "
        "The optimized MLP achieved 82.76% Accuracy, 75.00% Sensitivity (30/40 diabetic cases identified), 86.84% Specificity (66/76 non-diabetic cases identified), and 0.8681 ROC-AUC."
    )

    # -------------------------------------------------------------
    # SECTION 12: MODEL DEPLOYMENT AND TESTING
    # -------------------------------------------------------------
    add_h1("12 Model Deployment and Testing")
    add_body(
        "• 12.1 Deployment Approach: Production deployment using Python + Streamlit Community Cloud (zero React/JS dependencies).\n"
        "• 12.2 Deployment Tool: Streamlit 1.28+, Plotly 5.15+ for interactive gauges and radar charts, Joblib 1.3+ for model deserialization.\n"
        "• 12.3 User Interface: 'EndoPredict AI' multi-tab dashboard featuring single patient triage, biometric radar profile, interactive What-If sensitivity simulator, and batch cohort screening.\n"
        "• 12.4 Prediction Workflow: Raw patient input -> Validation -> Training-fitted median imputation -> Feature engineering -> Scaling -> MLPClassifier probability -> Decision threshold -> Risk stratum.\n"
        "• 12.5 Deployment Testing: Comprehensive automated test suite (test_deployment.py) passing 7/7 unit and contract tests in under 3.5 seconds."
    )

    # -------------------------------------------------------------
    # SECTION 13 TO 16: DISCUSSION, LIMITATIONS, CONCLUSION, FUTURE SCOPE
    # -------------------------------------------------------------
    add_h1("13 Results & Discussion")
    add_body(
        "The optimized MLP achieved strong predictive power by capturing non-linear metabolic risk surfaces. In clinical screening, "
        "False Negatives present high hazard because untreated diabetes causes irreversible organ damage. The dynamic threshold slider "
        "in EndoPredict AI allows clinicians to lower tau to 0.35, reducing False Negatives from 10 to 4."
    )

    add_h1("14 Limitations")
    add_body("The dataset is limited to N = 768 female patients of Pima Indian heritage; external clinical validation on broader multi-ethnic cohorts is required.")

    add_h1("15 Conclusion")
    add_body("A leak-free, reproducible machine learning system was successfully developed and deployed for diabetes onset prediction using an optimized Scikit-Learn Multilayer Perceptron.")

    add_h1("16 Future Scope")
    add_body("Future research includes multi-center cohort validation (NHANES), self-attention tabular architectures (TabNet), SHAP feature explainability, and mobile edge deployment.")

    # -------------------------------------------------------------
    # REFERENCES
    # -------------------------------------------------------------
    add_h1("References")
    refs_text = [
        "[1] J. W. Smith, J. E. Everhart, W. C. Dickson, W. C. Knowler, and R. S. Johannes, “Using the ADAP learning algorithm to forecast the onset of diabetes mellitus,” in Proc. Annu. Symp. Comput. Appl. Med. Care, 1988, pp. 261–265.",
        "[2] American Diabetes Association, “Standards of Medical Care in Diabetes—2024,” Diabetes Care, vol. 47, no. Suppl. 1, pp. S1–S343, 2024.",
        "[3] I. Kavakiotis et al., “Machine learning and data mining methods in diabetes research,” Comput. Struct. Biotechnol. J., vol. 15, pp. 104–116, 2017.",
        "[4] D. Sisodia and D. S. Sisodia, “Prediction of diabetes using classification algorithms,” Procedia Comput. Sci., vol. 132, pp. 1578–1585, 2018.",
        "[5] F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” J. Mach. Learn. Res., vol. 12, pp. 2825–2830, 2011."
    ]
    for r in refs_text:
        add_body(r)

    # -------------------------------------------------------------
    # APPENDICES
    # -------------------------------------------------------------
    add_h1("Appendix A: Source Code Structure")
    add_body(
        "The project source code is organized into modular single-responsibility scripts:\n"
        "• src/preprocessing.py: DiabetesPreprocessor pipeline handling zero detection, median imputation, and scaling.\n"
        "• src/feature_engineering.py: Mathematical implementation of all 16 domain features.\n"
        "• src/models.py: Factory functions for baseline classifiers and MLPClassifier.\n"
        "• src/train.py: Automated training, hyperparameter grid search, and model serialization.\n"
        "• src/prediction.py: Unified inference engine for single and batch predictions.\n"
        "• app/app.py: Enterprise Streamlit web dashboard."
    )

    add_h1("Appendix B: Additional Results")
    add_body("Full confusion matrices, precision-recall curves, and training cross-entropy loss trajectories are saved in the results/ directory.")

    add_h1("Appendix C: Deployment Screenshots")
    add_body("EndoPredict AI web interface is deployed at localhost:8501 and configured for Streamlit Community Cloud hosting.")

    add_h1("Appendix D: Contributions of Group Members")
    add_body("Individual student responsibilities across all project phases:")

    # Appendix D Table
    t_d = doc.add_table(rows=1, cols=5)
    t_d.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_d = t_d.rows[0].cells
    hdr_d_titles = ["Sl.", "Group Member", "Regd. No.", "Role / Responsibility", "Contribution (%)"]
    for i, title in enumerate(hdr_d_titles):
        hdr_d[i].paragraphs[0].text = title
        hdr_d[i].paragraphs[0].runs[0].font.bold = True
        hdr_d[i].paragraphs[0].runs[0].font.size = Pt(8.5)
        hdr_d[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_d[i], "0F172A")
        set_cell_margins(hdr_d[i], top=80, bottom=80, left=100, right=100)

    contrib_rows = [
        ["1", "Tribhuwan Singh", "2341019538", "ML Lead, MLP Architecture & Pipeline Architect", "25%"],
        ["2", "Surajit Sahoo", "2341019165", "Data Preprocessing, Imputation & Baseline Modeling", "25%"],
        ["3", "Anwesha Srichandan", "2341019594", "EDA, Feature Engineering & Hyperparameter Search", "25%"],
        ["4", "Priti Rani Maity", "2341013065", "Model Evaluation, Web Deployment & Report Preparation", "25%"]
    ]
    for row_data in contrib_rows:
        row = t_d.add_row().cells
        for i, val in enumerate(row_data):
            row[i].paragraphs[0].text = val
            row[i].paragraphs[0].runs[0].font.size = Pt(8.5)
            set_cell_margins(row[i], top=60, bottom=60, left=100, right=100)

    doc.save(DOCX_PATH)
    print(f"Laboratory Record DOCX successfully created at: {DOCX_PATH}")

if __name__ == "__main__":
    create_docx_report()
